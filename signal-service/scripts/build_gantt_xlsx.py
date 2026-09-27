"""Build an editable Excel Gantt for the CardioSaarthi schedule.

The bars are drawn by conditional formatting driven by the Start/End week cells,
so changing a number moves the bar. Nothing is painted in by hand -- a workbook
whose chart has to be re-coloured after every edit is not editable.
"""

from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(r"C:\Users\shado\Documents\PROJECTS26\CardioSaarthi\docs\planning\gantt_schedule.xlsx")

N_WEEKS = 16
START_MONDAY = date(2026, 7, 20)

PHASES = {
    1: ("ECG content preparation", "2F5FA6", "D1DCEB"),
    2: ("Student learning workflow", "3E8E5A", "D4E6DB"),
    3: ("AI-assisted learning & explanation", "E08A2E", "F8E5D1"),
    4: ("Faculty monitoring & validation", "7B5EA7", "E2DCEC"),
}

TASKS = [
    (1, "Dataset acquisition & quality screening", 1, 1, "Done"),
    (1, "Measurement engine (detection, delineation, intervals)", 1, 3, "Done"),
    (1, "Standards-compliant ECG rendering", 2, 3, "Done"),
    (1, "Technical validation vs LUDB & QTDB", 3, 3, "Done"),
    (1, "Clinical scenario generation (offline LLM batch)", 4, 5, "Not started"),
    (1, "Faculty review interface", 4, 5, "Not started"),
    (1, "Faculty review cycles -> approved case bank", 5, 8, "Not started"),
    (2, "Session state machine & step-lock orchestration", 7, 7, "Not started"),
    (2, "Deterministic grading + error taxonomy", 7, 8, "Not started"),
    (2, "Knowledge base ingestion & embeddings (pgvector)", 8, 8, "Not started"),
    (2, "Student interface - nine-step case player", 8, 10, "Not started"),
    (2, "Competency model & spaced repetition", 9, 10, "Not started"),
    (3, "LLM wrapper + CardioSaarthi persona configuration", 9, 9, "Not started"),
    (3, "RAG retrieval & citation pipeline", 9, 10, "Not started"),
    (3, "Output validation, fallback library, caching", 10, 11, "Not started"),
    (3, "Persona fixture suite (40 error scenarios)", 11, 11, "Not started"),
    (3, "Patient report explainer", 12, 13, "Not started"),
    (4, "Faculty analytics console", 11, 12, "Not started"),
    (4, "Technical validation report (final)", 13, 13, "Not started"),
    (4, "Study setup - ethics, consent, cohort scheduling", 12, 14, "Not started"),
    (4, "Pre-test / post-test with nursing cohort", 14, 15, "Not started"),
    (4, "Analysis, documentation & final report", 15, 16, "Not started"),
]

MILESTONES = [
    (3, "Engine validated against LUDB & QTDB", "Done"),
    (4, "First faculty review session", "Planned"),
    (8, "60 approved cases in the bank", "Planned"),
    (10, "Student runtime live end to end", "Planned"),
    (13, "Layers 3 and 4 complete", "Planned"),
    (16, "Final submission", "Planned"),
]

INK = "1A1A1A"
MUTED = "6B7680"
HAIRLINE = "D9DEE3"
INPUT_FILL = PatternFill("solid", fgColor="FFF6D5")

HDR_ROW = 8
FIRST_ROW = 9
WEEK_COL0 = 9  # column I

thin = Side(style="thin", color=HAIRLINE)
box = Border(left=thin, right=thin, top=thin, bottom=thin)


def base(sz=10, bold=False, colour=INK):
    return Font(name="Arial", size=sz, bold=bold, color=colour)


wb = Workbook()

# ---------------------------------------------------------------- Schedule
ws = wb.active
ws.title = "Schedule"
ws.sheet_view.showGridLines = False

ws["A1"] = "CardioSaarthi - Project Schedule"
ws["A1"].font = base(16, True)
ws["A2"] = "Sixteen weeks | Semester VII, 2026-27 | AI-ECG Tutor System"
ws["A2"].font = base(10, colour=MUTED)

ws["A4"] = "Project start (Monday of week 1)"
ws["C4"] = START_MONDAY
ws["C4"].fill = INPUT_FILL
ws["C4"].number_format = "dd mmm yyyy"
ws["A5"] = "Today"
ws["C5"] = "=TODAY()"
ws["C5"].number_format = "dd mmm yyyy"
ws["A6"] = "Current week"
ws["C6"] = f"=IFERROR(MAX(1,MIN({N_WEEKS},INT((C5-C4)/7)+1)),1)"
for r in (4, 5, 6):
    ws[f"A{r}"].font = base(10, colour=MUTED)
    ws[f"C{r}"].font = base(10, True)
ws["E4"] = "Yellow cells are inputs. Edit Start/End week and the bars move."
ws["E4"].font = base(9, colour=MUTED)

headers = ["Phase", "Task", "Start wk", "End wk", "Weeks", "Status", "Start date", "End date"]
for i, name in enumerate(headers, start=1):
    c = ws.cell(HDR_ROW, i, name)
    c.font = base(10, True)
    c.border = box
    c.alignment = Alignment(horizontal="left", vertical="center")
for w in range(1, N_WEEKS + 1):
    c = ws.cell(HDR_ROW, WEEK_COL0 + w - 1, w)
    c.font = base(9, True, MUTED)
    c.alignment = Alignment(horizontal="center", vertical="center")
    c.border = box

for i, (phase, name, start, end, status) in enumerate(TASKS):
    r = FIRST_ROW + i
    ws.cell(r, 1, f"Phase {phase}").font = base(10, colour=MUTED)
    ws.cell(r, 2, name).font = base(10)
    ws.cell(r, 3, start).fill = INPUT_FILL
    ws.cell(r, 4, end).fill = INPUT_FILL
    ws.cell(r, 5, f"=D{r}-C{r}+1")
    ws.cell(r, 6, status).fill = INPUT_FILL
    ws.cell(r, 7, f"=$C$4+(C{r}-1)*7")
    ws.cell(r, 8, f"=$C$4+D{r}*7-1")
    for col in (3, 4, 5, 6):
        ws.cell(r, col).font = base(10)
        ws.cell(r, col).alignment = Alignment(horizontal="center")
    for col in (7, 8):
        ws.cell(r, col).font = base(10, colour=MUTED)
        ws.cell(r, col).number_format = "dd mmm"
    for col in range(1, WEEK_COL0 + N_WEEKS):
        ws.cell(r, col).border = box

# Conditional formatting: one pair of rules per phase block. The "Done" rule is
# listed first so a completed task shows the solid colour rather than the pale.
first_week_letter = get_column_letter(WEEK_COL0)
last_week_letter = get_column_letter(WEEK_COL0 + N_WEEKS - 1)
row_of = {}
cursor = FIRST_ROW
for phase in PHASES:
    n = sum(1 for t in TASKS if t[0] == phase)
    row_of[phase] = (cursor, cursor + n - 1)
    cursor += n

def cf_fill(colour: str) -> PatternFill:
    """A solid fill that actually paints in a conditional-formatting rule.

    Conditional formats are stored as differential styles, and Excel takes the
    solid colour of a differential fill from bgColor, not fgColor. Setting
    fgColor alone -- which is what a normal cell fill uses -- produces a rule
    that evaluates correctly and paints nothing at all. Both are set here.
    """
    return PatternFill(fill_type="solid", start_color=colour, end_color=colour)


for phase, (r0, r1) in row_of.items():
    _label, solid, pale = PHASES[phase]
    rng = f"{first_week_letter}{r0}:{last_week_letter}{r1}"
    in_bar = f'AND({first_week_letter}${HDR_ROW}>=$C{r0},{first_week_letter}${HDR_ROW}<=$D{r0})'
    ws.conditional_formatting.add(
        rng,
        FormulaRule(formula=[f'AND({in_bar},$F{r0}="Done")'],
                    fill=cf_fill(solid), stopIfTrue=True),
    )
    ws.conditional_formatting.add(
        rng,
        FormulaRule(formula=[in_bar], fill=cf_fill(pale)),
    )

# Highlight the current week's column header.
ws.conditional_formatting.add(
    f"{first_week_letter}{HDR_ROW}:{last_week_letter}{HDR_ROW}",
    FormulaRule(formula=[f"{first_week_letter}${HDR_ROW}=$C$6"],
                fill=cf_fill("FCE1DE"),
                font=Font(name="Arial", size=9, bold=True, color="B3341F")),
)

dv = DataValidation(type="list", formula1='"Done,In progress,Not started"', allow_blank=True)
ws.add_data_validation(dv)
dv.add(f"F{FIRST_ROW}:F{FIRST_ROW + len(TASKS) - 1}")

# Phase rollup, computed from the task rows so it can never disagree with them.
sum_row = FIRST_ROW + len(TASKS) + 2
ws.cell(sum_row, 1, "PHASE SUMMARY").font = base(10, True)
for i, name in enumerate(["Phase", "Start wk", "End wk", "Tasks", "Done", "Complete"], start=1):
    c = ws.cell(sum_row + 1, i, name)
    c.font = base(10, True)
    c.border = box
last = FIRST_ROW + len(TASKS) - 1
for i, phase in enumerate(PHASES):
    r = sum_row + 2 + i
    key = f'"Phase {phase}"'
    ws.cell(r, 1, f"Phase {phase}").font = base(10)
    # SUMPRODUCT rather than MINIFS/MAXIFS: those need Excel 2019 or later and
    # yield #NAME? on anything older. SUMPRODUCT forces array evaluation in
    # every version without needing Ctrl+Shift+Enter. Non-matching rows are
    # pushed to 9999 for the minimum and to 0 for the maximum so they cannot win.
    a_rng = f"$A${FIRST_ROW}:$A${last}"
    ws.cell(r, 2, f"=SUMPRODUCT(MIN(({a_rng}={key})*$C${FIRST_ROW}:$C${last}"
                  f"+({a_rng}<>{key})*9999))")
    ws.cell(r, 3, f"=SUMPRODUCT(MAX(({a_rng}={key})*$D${FIRST_ROW}:$D${last}))")
    ws.cell(r, 4, f"=COUNTIFS($A${FIRST_ROW}:$A${last},{key})")
    ws.cell(r, 5, f"=COUNTIFS($A${FIRST_ROW}:$A${last},{key},$F${FIRST_ROW}:$F${last},\"Done\")")
    ws.cell(r, 6, f"=IFERROR(E{r}/D{r},0)")
    ws.cell(r, 6).number_format = "0%"
    for col in range(1, 7):
        ws.cell(r, col).border = box
        ws.cell(r, col).font = base(10)
        if col > 1:
            ws.cell(r, col).alignment = Alignment(horizontal="center")

note = sum_row + 2 + len(PHASES) + 1
ws.cell(note, 1, "Critical path: faculty review throughput in weeks 5-8, not engineering.")
ws.cell(note, 1).font = base(9, colour=MUTED)

ws.column_dimensions["A"].width = 10
ws.column_dimensions["B"].width = 48
# C holds the project start and today's date in "dd mmm yyyy"; 11 is too narrow
# for that and Excel renders ########.
ws.column_dimensions["C"].width = 14
for col in "DEFGH":
    ws.column_dimensions[col].width = 11
for w in range(N_WEEKS):
    ws.column_dimensions[get_column_letter(WEEK_COL0 + w)].width = 4.3
ws.freeze_panes = f"C{FIRST_ROW}"

# -------------------------------------------------------------- Milestones
ms = wb.create_sheet("Milestones")
ms.sheet_view.showGridLines = False
ms["A1"] = "Milestones"
ms["A1"].font = base(14, True)
for i, name in enumerate(["Week", "Milestone", "Status", "Date"], start=1):
    c = ms.cell(3, i, name)
    c.font = base(10, True)
    c.border = box
for i, (week, label, status) in enumerate(MILESTONES):
    r = 4 + i
    ms.cell(r, 1, week).fill = INPUT_FILL
    ms.cell(r, 2, label)
    ms.cell(r, 3, status).fill = INPUT_FILL
    ms.cell(r, 4, f"=Schedule!$C$4+A{r}*7-1")
    ms.cell(r, 4).number_format = "dd mmm yyyy"
    for col in range(1, 5):
        ms.cell(r, col).border = box
        ms.cell(r, col).font = base(10)
    ms.cell(r, 1).alignment = Alignment(horizontal="center")
ms.column_dimensions["A"].width = 8
ms.column_dimensions["B"].width = 44
ms.column_dimensions["C"].width = 14
ms.column_dimensions["D"].width = 16

# ----------------------------------------------------------------- Read me
rm = wb.create_sheet("Read me")
rm.sheet_view.showGridLines = False
lines = [
    ("CardioSaarthi - Project Schedule", 14, True, INK),
    ("", 10, False, INK),
    ("How to edit", 11, True, INK),
    ("Yellow cells are the only ones to type in:", 10, False, INK),
    ("   Schedule!C4       project start date (Monday of week 1)", 10, False, MUTED),
    ("   Start wk / End wk  week numbers for each task", 10, False, MUTED),
    ("   Status             Done / In progress / Not started (dropdown)", 10, False, MUTED),
    ("", 10, False, INK),
    ("Everything else is a formula. The Gantt bars are conditional formatting", 10, False, MUTED),
    ("driven by Start wk and End wk, so editing a week number moves the bar", 10, False, MUTED),
    ("immediately - there is no chart to redraw and no colour to repaint.", 10, False, MUTED),
    ("Marking a task Done turns its bar from pale to solid.", 10, False, MUTED),
    ("", 10, False, INK),
    ("Colour", 11, True, INK),
    ("Solid bar    delivered and validated", 10, False, MUTED),
    ("Pale bar     planned", 10, False, MUTED),
    ("Phase colours match the methodology diagram, so the slides read as one system.", 10, False, MUTED),
    ("", 10, False, INK),
    ("Status at end of week 3", 11, True, INK),
    ("Phase 1 is 4 of 7 tasks complete. The measurement engine, the renderer and", 10, False, MUTED),
    ("the technical validation are delivered. The case bank is not, because a case", 10, False, MUTED),
    ("bank is not finished until faculty have approved it.", 10, False, MUTED),
    ("", 10, False, INK),
    ("Critical path", 11, True, INK),
    ("Faculty review throughput in weeks 5-8, and study setup in weeks 12-14.", 10, False, MUTED),
    ("Neither is a coding task. Both should be booked in advance rather than", 10, False, MUTED),
    ("requested when needed.", 10, False, MUTED),
]
for i, (text, size, bold, colour) in enumerate(lines, start=1):
    c = rm.cell(i, 1, text)
    c.font = base(size, bold, colour)
rm.column_dimensions["A"].width = 92

OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print(f"written -> {OUT}")
